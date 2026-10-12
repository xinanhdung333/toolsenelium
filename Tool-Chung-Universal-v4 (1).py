"""
TOOL CHUNG CHO TẤT CẢ LOẠI WEB FORM  (bản đã sửa lỗi)

[v4] Giao diện mới: header, thẻ (card), nút màu, thống kê PASS/FAIL trực tiếp, nhật ký nền tối.

- Quét form để phát hiện field (input, select, textarea, radio, checkbox)
- Map cột Excel -> field web (tự động + sửa tay, hiện đủ mọi cột Excel)
- Hỗ trợ ngày dạng TEXT trong Excel: dd/mm/yyyy, mm/dd/yyyy, yyyy-mm-dd...
- Bắt lỗi đỏ, thông báo thành công, alert() và lỗi validation HTML5 của trình duyệt
- Chấm PASS / FAIL / REVIEW (không bắt được gì -> cần xem lại, KHÔNG tự cho PASS)
- Nhận file HTML local hoặc URL http(s)

[v2] Thêm: nút "🖱 Chọn" - mở Chrome, GIỮ CTRL + CLICK vào nút submit / vùng lỗi / vùng thành công
để tool dùng đúng phần tử đó thay vì tự đoán.

Cách dùng:
1. Chọn file HTML (hoặc dán URL) và Excel
2. "Quét Form Web" -> "Load Excel" -> "Tự Động Map" (double-click để sửa)
3. Chọn format ngày, tên cột Expected, (tuỳ chọn) CSS nút submit và từ khoá thành công
4. "CHẠY TEST"

Cột Expected hiểu như sau:
- "Thành công" / "Success" / "Pass"            -> mong đợi thành công
- "Không thành công", "Fail", "Lỗi", "Error"   -> mong đợi lỗi (không kiểm nội dung)
- "Email không hợp lệ", "Vui lòng nhập tên"... -> mong đợi lỗi VÀ nội dung lỗi phải khớp
- "Fail - Email không hợp lệ"                   -> mong đợi lỗi + kiểm nội dung "Email không hợp lệ"
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
import os
import re
import time
import html as html_lib
import threading
import unicodedata
import webbrowser
from pathlib import Path
from datetime import datetime, date, timedelta

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.common.exceptions import NoAlertPresentException, NoSuchElementException

# ----------------------------------------------------------------------------
# Hằng số & hàm tiện ích
# ----------------------------------------------------------------------------

DATE_FORMATS = {
    "dd/mm/yyyy": "%d/%m/%Y",
    "mm/dd/yyyy": "%m/%d/%Y",
    "yyyy/mm/dd": "%Y/%m/%d",
    "yyyy/dd/mm": "%Y/%d/%m",
    "dd-mm-yyyy": "%d-%m-%Y",
    "mm-dd-yyyy": "%m-%d-%Y",
    "yyyy-mm-dd": "%Y-%m-%d",
    "dd/mm/yy": "%d/%m/%y",
    "Text giữ nguyên": "TEXT",
}

PLACEHOLDERS = {"(trống)", "(trong)", "null", "none", "-"}
TRUE_SET = {"1", "true", "yes", "y", "x", "co", "check", "checked", "on", "v"}

# Từ khoá (đã bỏ dấu, chữ thường) coi là THẤT BẠI. Luôn kiểm tra trước từ khoá thành công
# để "không thành công" không bị hiểu nhầm là "thành công".
FAIL_KWS = [
    "loi", "that bai", "khong thanh cong", "fail", "failed", "error", "invalid",
    "khong hop le", "unsuccessful", "sai", "khong duoc", "vui long", "bat buoc",
    "khong de trong", "required",
]
NEG_EXPECT = [
    "khong thanh cong", "that bai", "fail", "failed", "loi", "error", "invalid",
    "khong hop le", "unsuccessful", "sai", "khong duoc", "bao loi",
]
LEADING_GENERIC = r"^(?:fail|failed|error|loi|that bai|khong thanh cong|unsuccessful|invalid)\b[\s\-:]*"

JS_SET = ("arguments[0].value=arguments[1];"
          "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
          "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));")

JS_LABEL = """
var el = arguments[0], t = '';
if (el.id) { var l = document.querySelector('label[for="' + CSS.escape(el.id) + '"]'); if (l) t = l.innerText; }
if (!t) { var c = el.closest('label'); if (c) t = c.innerText; }
if (!t) { var p = el.previousElementSibling; if (p && p.tagName == 'LABEL') t = p.innerText; }
if (!t && el.parentElement) { var l2 = el.parentElement.querySelector('label'); if (l2) t = l2.innerText; }
if (!t) t = el.getAttribute('aria-label') || el.getAttribute('placeholder') || '';
return (t || '').trim();
"""

JS_INVALID = """
return Array.from(document.querySelectorAll('input:invalid, select:invalid, textarea:invalid'))
  .map(function(e){ return (e.id || e.name || e.tagName) + ': ' + e.validationMessage; });
"""

# ---- Bảng màu & font giao diện hiện đại
BG, CARD, LINE = "#f1f5f9", "#ffffff", "#e2e8f0"
TEXT, MUTED = "#0f172a", "#64748b"
PRIMARY, PRIMARY_DARK = "#2563eb", "#1e3a8a"
FONT = ("Segoe UI", 10)
FONT_B = ("Segoe UI", 10, "bold")
STAT_COLORS = {   # nền, chữ
    "TOTAL": ("#e0e7ff", "#3730a3"), "PASS": ("#dcfce7", "#166534"), "FAIL": ("#fee2e2", "#991b1b"),
    "REVIEW": ("#fef3c7", "#92400e"), "ERROR": ("#ede9fe", "#5b21b6"),
}
STATUS_ICON = {"PASS": "✔ PASS", "FAIL": "✖ FAIL", "REVIEW": "⚠ REVIEW", "ERROR": "‼ ERROR"}

REPORT_TEMPLATE = r'''<!DOCTYPE html>
<html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Báo cáo kiểm thử tự động</title>
<style>
:root{--bg:#f1f5f9;--card:#fff;--text:#0f172a;--muted:#64748b;--line:#e2e8f0;--head:#1e293b;
--pass:#16a34a;--passbg:#dcfce7;--fail:#dc2626;--failbg:#fee2e2;--review:#d97706;--reviewbg:#fef3c7;--error:#7c3aed;--errorbg:#ede9fe;--accent:#2563eb}
@media (prefers-color-scheme:dark){:root{--bg:#0b1220;--card:#111a2e;--text:#e2e8f0;--muted:#94a3b8;--line:#223052;--head:#0f172a;
--passbg:#14532d55;--failbg:#7f1d1d55;--reviewbg:#78350f55;--errorbg:#4c1d9555}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--text);font-family:Inter,"Segoe UI",system-ui,sans-serif;font-size:14px}
header{background:linear-gradient(120deg,#1e3a8a,#2563eb 60%,#06b6d4);color:#fff;padding:28px 32px}
header h1{margin:0 0 6px;font-size:24px}
header p{margin:2px 0;opacity:.9;font-size:13px;word-break:break-all}
main{max-width:1400px;margin:-22px auto 40px;padding:0 20px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:14px}
.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:16px 18px;box-shadow:0 2px 8px rgba(15,23,42,.06)}
.card .n{font-size:32px;font-weight:800;line-height:1.1}
.card .l{color:var(--muted);font-size:12px;text-transform:uppercase;letter-spacing:.05em;margin-top:4px}
.c-pass{border-top:4px solid var(--pass)}.c-pass .n{color:var(--pass)}
.c-fail{border-top:4px solid var(--fail)}.c-fail .n{color:var(--fail)}
.c-review{border-top:4px solid var(--review)}.c-review .n{color:var(--review)}
.c-error{border-top:4px solid var(--error)}.c-error .n{color:var(--error)}
.c-total{border-top:4px solid var(--accent)}
.overview{display:grid;grid-template-columns:260px 1fr;gap:14px;margin-top:14px}
@media(max-width:800px){.overview{grid-template-columns:1fr}}
.donutwrap{display:flex;align-items:center;justify-content:center;gap:16px}
.donut{width:140px;height:140px;border-radius:50%;position:relative}
.donut::after{content:"";position:absolute;inset:22px;background:var(--card);border-radius:50%}
.donut b{position:absolute;inset:0;display:flex;align-items:center;justify-content:center;z-index:1;font-size:26px}
.kv{display:grid;grid-template-columns:130px 1fr;gap:6px 12px;font-size:13px}
.kv div:nth-child(odd){color:var(--muted)}
.kv div:nth-child(even){word-break:break-all}
h2{font-size:16px;margin:26px 0 10px}
.alerts{display:grid;grid-template-columns:repeat(auto-fill,minmax(330px,1fr));gap:12px}
.alert{background:var(--card);border:1px solid var(--line);border-left:5px solid var(--fail);border-radius:10px;padding:12px 14px;cursor:pointer}
.alert.REVIEW{border-left-color:var(--review)}.alert.ERROR{border-left-color:var(--error)}
.alert .t{font-weight:700;display:flex;justify-content:space-between;gap:8px}
.alert .d{color:var(--muted);font-size:12px;margin:3px 0}
.alert .r{font-size:13px}
.toolbar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin-bottom:10px}
.chipbtn{border:1px solid var(--line);background:var(--card);color:var(--text);padding:6px 14px;border-radius:999px;cursor:pointer;font-size:13px}
.chipbtn.active{background:var(--accent);border-color:var(--accent);color:#fff}
.toolbar input{flex:1;min-width:200px;padding:8px 12px;border:1px solid var(--line);border-radius:8px;background:var(--card);color:var(--text)}
.toolbar button.act{padding:7px 12px;border:1px solid var(--line);background:var(--card);color:var(--text);border-radius:8px;cursor:pointer}
.tablewrap{background:var(--card);border:1px solid var(--line);border-radius:14px;overflow:auto;box-shadow:0 2px 8px rgba(15,23,42,.06)}
table{border-collapse:collapse;width:100%}
th{background:var(--head);color:#fff;text-align:left;padding:10px 12px;font-size:12px;text-transform:uppercase;letter-spacing:.04em;position:sticky;top:0;z-index:2}
td{padding:10px 12px;border-top:1px solid var(--line);vertical-align:top}
tr.FAIL td:first-child{box-shadow:inset 4px 0 var(--fail)}
tr.REVIEW td:first-child{box-shadow:inset 4px 0 var(--review)}
tr.ERROR td:first-child{box-shadow:inset 4px 0 var(--error)}
tr.PASS td:first-child{box-shadow:inset 4px 0 var(--pass)}
tr:hover td{background:rgba(37,99,235,.05)}
.tc{font-weight:700;white-space:nowrap}
.sc{color:var(--muted);font-size:12px;max-width:260px;margin-top:2px}
.chip{display:inline-block;background:rgba(100,116,139,.12);border-radius:6px;padding:2px 7px;margin:2px 4px 2px 0;font-size:12px;max-width:260px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;vertical-align:top}
.chip b{color:var(--muted);font-weight:600;margin-right:4px}
.chip i{color:var(--muted)}
.badge{display:inline-block;border-radius:999px;padding:3px 11px;font-weight:700;font-size:12px;white-space:nowrap}
.b-PASS,.b-SUCCESS{background:var(--passbg);color:var(--pass)}
.b-FAIL{background:var(--failbg);color:var(--fail)}
.b-REVIEW,.b-UNKNOWN{background:var(--reviewbg);color:var(--review)}
.b-ERROR{background:var(--errorbg);color:var(--error)}
.b-neutral{background:rgba(100,116,139,.15);color:var(--muted)}
.errtxt{color:var(--fail);font-size:12.5px}
.oktxt{color:var(--pass);font-size:12.5px}
.note{color:var(--muted);font-size:12.5px}
.thumb{width:84px;border:1px solid var(--line);border-radius:6px;cursor:zoom-in;display:block}
.modal{position:fixed;inset:0;background:rgba(2,6,23,.8);display:none;align-items:center;justify-content:center;z-index:99;padding:20px}
.modal.on{display:flex}.modal img{max-width:95%;max-height:92vh;border-radius:8px;background:#fff}
details{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:10px 14px;margin-top:14px}
summary{cursor:pointer;font-weight:600}
footer{text-align:center;color:var(--muted);font-size:12px;margin:24px 0}
.empty{padding:30px;text-align:center;color:var(--muted)}
@media print{header{background:#1e3a8a!important;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.toolbar,.modal{display:none!important}.tablewrap{overflow:visible}th{position:static}}
</style></head><body>
<header>
  <h1>🧪 Báo cáo kiểm thử tự động</h1>
  <p id="h-web"></p><p id="h-time"></p>
</header>
<main>
  <div class="cards" id="cards"></div>
  <div class="overview">
    <div class="card donutwrap"><div class="donut" id="donut"><b id="rate"></b></div></div>
    <div class="card"><div class="kv" id="kv"></div></div>
  </div>
  <div id="attention"></div>
  <h2>Chi tiết từng test case</h2>
  <div class="toolbar" id="toolbar"></div>
  <div class="tablewrap"><table><thead><tr>
    <th>#</th><th>Test case</th><th>Dữ liệu nhập</th><th>Expected</th><th>Actual</th>
    <th>Thông báo thực tế</th><th>Ghi chú</th><th>Kết quả</th><th>Ảnh</th></tr></thead>
    <tbody id="rows"></tbody></table>
    <div class="empty" id="empty" style="display:none">Không có test case nào khớp bộ lọc.</div></div>
  <details><summary>Cấu hình lần chạy</summary><div class="kv" id="cfg" style="margin-top:10px"></div></details>
  <footer>Tạo tự động bởi Tool Chung Universal</footer>
</main>
<div class="modal" id="modal"><img id="modalimg" alt=""></div>
<script>
const META=__META__;
const R=__RESULTS__;
const $=id=>document.getElementById(id);
const esc=s=>String(s==null?"":s).replace(/[&<>"']/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
const COL={PASS:"#16a34a",FAIL:"#dc2626",REVIEW:"#d97706",ERROR:"#7c3aed"};
const C=META.counts, T=META.total||1;
$("h-web").textContent="Form: "+META.web;
$("h-time").textContent="Thời gian chạy: "+META.time+"  •  Tổng thời lượng: "+META.seconds+" giây";
$("cards").innerHTML=
 `<div class="card c-total"><div class="n">${META.total}</div><div class="l">Tổng test case</div></div>`+
 ["PASS","FAIL","REVIEW","ERROR"].map(k=>`<div class="card c-${k.toLowerCase()}"><div class="n">${C[k]}</div><div class="l">${k}</div></div>`).join("");
let acc=0,segs=[];
["PASS","FAIL","REVIEW","ERROR"].forEach(k=>{const p=C[k]/T*100;if(p>0){segs.push(`${COL[k]} ${acc}% ${acc+p}%`);acc+=p;}});
$("donut").style.background="conic-gradient("+(segs.join(",")||"#e5e7eb 0 100%")+")";
$("rate").textContent=Math.round(C.PASS/T*100)+"%";
const kv=(a)=>a.map(([k,v])=>`<div>${esc(k)}</div><div>${v}</div>`).join("");
const avg=R.length?(R.reduce((s,r)=>s+(r.duration||0),0)/R.length).toFixed(2):"0";
$("kv").innerHTML=kv([
 ["Tỉ lệ đạt (PASS)",`<b>${Math.round(C.PASS/T*1000)/10}%</b> (${C.PASS}/${META.total})`],
 ["Phát hiện lỗi form",`<b style="color:var(--fail)">${C.FAIL}</b> test case FAIL`],
 ["Cần xem tay",`${C.REVIEW} REVIEW, ${C.ERROR} ERROR`],
 ["Thời gian TB / case",avg+" giây"],
 ["Số cột dữ liệu nhập",META.mapping.length]]);
$("cfg").innerHTML=kv([
 ["File form",esc(META.web)],["File Excel",esc(META.excel)],["Format ngày",esc(META.date_format)],
 ["Map cột → field",META.mapping.map(m=>`<span class="chip"><b>${esc(m[0])}</b>→ ${esc(m[1])}</span>`).join("")],
 ["Nút submit",esc(META.selectors.submit)],["Vùng báo lỗi",esc(META.selectors.error)],["Vùng thành công",esc(META.selectors.msg)],
 ["Ý nghĩa trạng thái","PASS: khớp Expected • FAIL: lệch Expected (lỗi của form) • REVIEW: tool không bắt được thông báo, cần xem tay • ERROR: tool gặp sự cố khi chạy"]]);
// ---- cần chú ý
const bad=R.filter(r=>r.status!=="PASS");
if(bad.length){
 $("attention").innerHTML="<h2>⚠️ Cần chú ý ("+bad.length+")</h2><div class='alerts'>"+bad.map(r=>
  `<div class="alert ${r.status}" onclick="jump('${esc(r.tc_id)}')"><div class="t"><span>${esc(r.tc_id)}</span><span class="badge b-${r.status}">${r.status}</span></div>
   <div class="d">${esc(r.scenario||"")}</div>
   <div class="r">${esc(r.note||r.error||"")}</div></div>`).join("")+"</div>";
}
// ---- bảng
let filter="ALL",q="";
const tb=$("toolbar");
tb.innerHTML=["ALL","PASS","FAIL","REVIEW","ERROR"].map(k=>`<button class="chipbtn${k==="ALL"?" active":""}" data-k="${k}">${k==="ALL"?"Tất cả ("+META.total+")":k+" ("+C[k]+")"}</button>`).join("")+
 `<input id="q" placeholder="Tìm theo mã, mô tả, dữ liệu, thông báo..."><button class="act" onclick="exportCSV()">⬇ CSV</button><button class="act" onclick="window.print()">🖨 In</button>`;
tb.onclick=e=>{const k=e.target.dataset&&e.target.dataset.k;if(!k)return;filter=k;[...tb.querySelectorAll(".chipbtn")].forEach(b=>b.classList.toggle("active",b.dataset.k===k));render();};
$("q").oninput=e=>{q=e.target.value.toLowerCase();render();};
function dataChips(d){return (d||[]).map(([k,v])=>{const e=(String(v).trim()===""||v==="(trống)");return `<span class="chip" title="${esc(k)}: ${esc(v)}"><b>${esc(k.split(" (")[0])}</b>${e?"<i>(trống)</i>":esc(v)}</span>`}).join("");}
function expBadge(e){const n=(e||"").toLowerCase();const ok=/thành công|thanh cong|success|pass/.test(n)&&!/không|khong|fail/.test(n);
 return `<span class="badge ${e?(ok?"b-SUCCESS":"b-FAIL"):"b-neutral"}">${esc(e||"—")}</span>`;}
function render(){
 const rows=R.map((r,i)=>[r,i]).filter(([r])=>{
  if(filter!=="ALL"&&r.status!==filter)return false;
  if(!q)return true;
  return (r.tc_id+" "+(r.scenario||"")+" "+(r.data||[]).map(x=>x.join(" ")).join(" ")+" "+(r.error||"")+" "+(r.note||"")).toLowerCase().includes(q);});
 $("rows").innerHTML=rows.map(([r,i])=>`<tr class="${r.status}" id="tc-${esc(r.tc_id)}">
  <td>${i+1}</td>
  <td><div class="tc">${esc(r.tc_id)}</div><div class="sc">${esc(r.scenario||"")}</div><div class="sc">⏱ ${r.duration||0}s</div></td>
  <td>${dataChips(r.data)}</td>
  <td>${expBadge(r.expected)}</td>
  <td><span class="badge b-${esc(r.actual)}">${esc(r.actual)}</span></td>
  <td class="${r.actual==="SUCCESS"?"oktxt":"errtxt"}">${esc(r.error||"")}</td>
  <td class="note">${esc(r.note||"")}</td>
  <td><span class="badge b-${r.status}">${r.status}</span></td>
  <td>${r.shot?`<img class="thumb" src="data:image/png;base64,${r.shot}" onclick="zoom(this.src)">`:""}</td></tr>`).join("");
 $("empty").style.display=rows.length?"none":"block";
}
function jump(id){filter="ALL";q="";$("q").value="";[...tb.querySelectorAll(".chipbtn")].forEach(b=>b.classList.toggle("active",b.dataset.k==="ALL"));render();
 const el=document.getElementById("tc-"+id);if(el){el.scrollIntoView({behavior:"smooth",block:"center"});el.style.outline="3px solid #2563eb";setTimeout(()=>el.style.outline="",1800);}}
function zoom(src){$("modalimg").src=src;$("modal").classList.add("on");}
$("modal").onclick=()=>$("modal").classList.remove("on");
function exportCSV(){
 const q2=s=>'"'+String(s==null?"":s).replace(/"/g,'""')+'"';
 const lines=[["TC_ID","Mô tả","Dữ liệu","Expected","Actual","Thông báo thực tế","Ghi chú","Kết quả"].map(q2).join(",")];
 R.forEach(r=>lines.push([r.tc_id,r.scenario,(r.data||[]).map(x=>x[0]+"="+x[1]).join(" | "),r.expected,r.actual,r.error,r.note,r.status].map(q2).join(",")));
 const a=document.createElement("a");a.href=URL.createObjectURL(new Blob(["\ufeff"+lines.join("\n")],{type:"text/csv;charset=utf-8"}));
 a.download="ket_qua_kiem_thu.csv";a.click();}
render();
</script></body></html>
'''

JS_PICK = """
var group = arguments[0];
window.__pickGroup = group;
if (window.__pickerReady) { return; }
window.__pickerReady = true; window.__picked = null;
function sel(el) {
  var path = [];
  while (el && el.nodeType == 1 && el !== document.body) {
    if (el.id) { path.unshift('#' + CSS.escape(el.id)); break; }
    var s = el.tagName.toLowerCase();
    var cls = Array.from(el.classList).filter(function(c){ return !/^[0-9]/.test(c); }).slice(0, 2)
                   .map(function(c){ return '.' + CSS.escape(c); }).join('');
    s += cls;
    var sib = Array.from(el.parentNode.children).filter(function(x){ return x.tagName === el.tagName; });
    if (sib.length > 1) s += ':nth-of-type(' + (sib.indexOf(el) + 1) + ')';
    path.unshift(s); el = el.parentElement;
  }
  return path.join(' > ');
}
function groupSel(el) {
  var cl = Array.from(el.classList);
  for (var i = 0; i < cl.length; i++) {
    try { if (document.querySelectorAll('.' + CSS.escape(cl[i])).length > 1) return '.' + CSS.escape(cl[i]); } catch (e) {}
  }
  var r = el.getAttribute('role');
  if (r) return '[role="' + r + '"]';
  return null;
}
var bar = document.createElement('div');
bar.textContent = 'TOOL: thao tác bình thường, GIỮ CTRL + CLICK vào phần tử cần chọn';
bar.style.cssText = 'position:fixed;top:0;left:0;right:0;background:#f97316;color:#fff;font:bold 14px Arial;padding:6px;text-align:center;z-index:2147483647;pointer-events:none';
document.body.appendChild(bar);
var last = null;
document.addEventListener('mouseover', function(e) {
  if (last) last.style.outline = last.__old || '';
  last = e.target; last.__old = last.style.outline; last.style.outline = '2px solid #f97316';
}, true);
document.addEventListener('click', function(e) {
  if (!(e.ctrlKey || e.metaKey)) return;
  e.preventDefault(); e.stopPropagation();
  var s = window.__pickGroup ? groupSel(e.target) : null;
  window.__picked = s || sel(e.target);
}, true);
"""

JS_ONLY_TYPES = {"time", "month", "week", "datetime-local", "color", "range"}

ERROR_SELECTOR = "[id*='error'], .error, [class*='error'], [role='alert']"
MSG_SELECTOR = "#message, .success, [class*='success'], [id*='success'], [role='status']"


def norm(s):
    """Bỏ dấu tiếng Việt, chữ thường, chỉ giữ chữ-số, gộp khoảng trắng."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return s.strip()


def kw_hit(text_norm, kw_norm):
    """Khớp theo ranh giới từ (tránh 'sai' khớp trong 'saigon')."""
    if not kw_norm:
        return False
    return f" {kw_norm} " in f" {text_norm} "


def to_url(src):
    s = src.strip()
    if re.match(r"^(https?|file):", s, re.I):
        return s
    return Path(s).resolve().as_uri()


def source_exists(src):
    s = src.strip()
    if not s:
        return False
    if re.match(r"^https?:", s, re.I):
        return True
    if s.lower().startswith("file:"):
        return True
    return os.path.exists(s)


def clean_cell(v):
    """Chuẩn hoá giá trị ô Excel. Trả '' / str / datetime / date."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.lower() in PLACEHOLDERS:
        return ""
    return s


def to_text(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return str(v)


def parse_date_flexible(text, fmt_key):
    """Trả (iso 'yyyy-mm-dd' hoặc '', ghi chú). Ngày sai (vd 31/02) -> ('', cảnh báo)."""
    if text is None or text == "":
        return "", ""
    if isinstance(text, datetime):
        return text.strftime("%Y-%m-%d"), ""
    if isinstance(text, date):
        return text.strftime("%Y-%m-%d"), ""
    s = str(text).strip()
    if not s or s.lower() in PLACEHOLDERS:
        return "", ""

    # Ngày tương đối
    low = norm(s)
    today = date.today()
    if low in ("hom nay", "today"):
        return today.isoformat(), ""
    if low in ("hom qua", "yesterday"):
        return (today - timedelta(days=1)).isoformat(), ""
    if low in ("ngay mai", "tomorrow"):
        return (today + timedelta(days=1)).isoformat(), ""
    m = re.match(r"^(?:hôm nay|today)?\s*([+-])\s*(\d+)\s*(?:ngày|days?|d)?$", s.lower())
    if m:
        n = int(m.group(2))
        return (today + timedelta(days=n if m.group(1) == "+" else -n)).isoformat(), ""

    m = re.match(r"^(\d{1,4})[/\-.](\d{1,2})[/\-.](\d{1,4})$", s)
    if not m:
        return "", f"Không nhận dạng được ngày '{s}'"
    g = m.groups()
    note = ""

    if fmt_key == "Text giữ nguyên":
        if len(g[0]) == 4:                      # yyyy-?-?
            order = ["yyyy", "mm", "dd"] if int(g[1]) <= 12 else ["yyyy", "dd", "mm"]
        elif int(g[1]) > 12:                    # giữa > 12 -> chắc chắn mm/dd
            order = ["mm", "dd", "yyyy"]
        else:
            order = ["dd", "mm", "yyyy"]
            if int(g[0]) <= 12 and g[0] != g[1]:
                note = f"Ngày '{s}' mơ hồ, tool đoán dd/mm"
    else:
        order = re.findall(r"yyyy|yy|dd|mm", fmt_key.lower())
        if len(order) != 3:
            order = ["dd", "mm", "yyyy"]

    parts = {}
    for name, val in zip(order, g):
        key = "yyyy" if name in ("yyyy", "yy") else name
        parts[key] = val
    y = parts.get("yyyy", "")
    mo = parts.get("mm", "")
    d = parts.get("dd", "")
    try:
        yi = int(y)
        if len(y) <= 2:
            yi += 2000
        return date(yi, int(mo), int(d)).isoformat(), note
    except (ValueError, TypeError):
        return "", f"Ngày không hợp lệ '{s}' (theo format {fmt_key}) -> để trống"


def parse_expected(text, success_kws):
    """Trả (should_pass, expected_message_norm). should_pass=None nếu Expected trống."""
    n = norm(text)
    if not n:
        return None, ""
    for kw in NEG_EXPECT:
        if kw_hit(n, kw):
            break
    else:
        if any(kw_hit(n, k) for k in list(success_kws) + ["pass", "passed"]):
            return True, ""
    # Còn lại: coi như mong đợi lỗi; phần sau chữ Fail/Lỗi/... là nội dung lỗi cần khớp
    msg = re.sub(LEADING_GENERIC, "", n).strip()
    if msg in ("", "khong thanh cong") or len(msg) < 3:
        msg = ""
    return False, msg


def msg_matches(exp_msg_norm, actual_text):
    en = norm(actual_text)
    if exp_msg_norm in en:
        return True
    toks = exp_msg_norm.split()
    if not toks:
        return True
    have = set(en.split())
    return sum(1 for t in toks if t in have) / len(toks) >= 0.7


# Nhóm từ đồng nghĩa dùng cho Tự Động Map (không gắn cứng web đặt vé nữa)
SYNONYM_GROUPS = [
    # (từ khoá ở header (đã bỏ dấu), từ khoá ở key/label field, type input, tag)
    (["ngay", "date", "checkin", "checkout", "check in", "check out"], ["date", "ngay"], {"date"}, set()),
    (["email", "mail"], ["email", "mail"], {"email"}, set()),
    (["sdt", "phone", "dien thoai", "tel", "mobile"], ["phone", "tel", "sdt", "mobile", "dienthoai"], {"tel"}, set()),
    (["password", "mat khau", "matkhau", "mk"], ["password", "matkhau", "pass"], {"password"}, set()),
    (["phim", "movie", "ban", "khoa", "course", "loai", "type", "room", "phong"],
     ["movie", "course", "room", "type", "phim", "table", "khoa", "loai"], set(), {"select"}),
    (["so luong", "ve", "ticket", "nguoi", "guest", "quantity", "qty", "so"],
     ["ticket", "quantity", "qty", "guest", "people", "amount", "soluong"], {"number"}, set()),
    (["ten", "name", "ho ten", "hoten"], ["name", "ten", "hoten", "fullname"], set(), set()),
]

NON_INPUT_HEADERS = {
    "tc id", "tcid", "tc", "test scenario", "test case", "expected", "expected result",
    "should pass", "ghi chu", "note", "notes", "mo ta", "description", "actual",
    "status", "ket qua", "result", "stt", "no", "priority", "do uu tien",
}


# ----------------------------------------------------------------------------
# GUI
# ----------------------------------------------------------------------------

class UniversalTesterGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("TOOL CHUNG - Test Mọi Loại Web Form")
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        w, h = min(1360, sw - 40), min(900, sh - 90)          # không vượt quá màn hình
        self.root.geometry(f"{w}x{h}+{max(0, (sw - w) // 2)}+10")
        self.test_data = []
        self.results = []
        self.web_fields = []
        self.excel_headers = []
        self.mappings = {}          # excel_col_index -> field key
        self.running = False
        self.stop_event = threading.Event()
        self.report_path = ""

        self.setup_style()
        root.configure(bg=BG)
        root.minsize(900, 540)

        # ---------- HEADER ----------
        header = tk.Frame(root, bg=PRIMARY_DARK)
        header.pack(fill="x")
        tk.Label(header, text="🧪  TOOL CHUNG  ·  Test Mọi Loại Web Form", bg=PRIMARY_DARK, fg="white",
                 font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=22, pady=(14, 0))
        tk.Label(header, text="Quét form  →  Load Excel  →  Tự động map  →  Chạy test  →  Xem báo cáo",
                 bg=PRIMARY_DARK, fg="#bfdbfe", font=("Segoe UI", 10)).pack(anchor="w", padx=22, pady=(2, 14))

        # Thanh cố định (luôn hiện): nút chạy + thống kê + tiến trình
        top_fixed = ttk.Frame(root, padding=(16, 10, 16, 0))
        top_fixed.pack(fill="x")

        # Vùng nội dung CUỘN ĐƯỢC (con lăn chuột hoặc thanh cuộn bên phải)
        container = ttk.Frame(root)
        container.pack(fill="both", expand=True)
        self.canvas = tk.Canvas(container, bg=BG, highlightthickness=0)
        vsb = ttk.Scrollbar(container, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        body = ttk.Frame(self.canvas, padding=(16, 8, 16, 12))
        win = self.canvas.create_window((0, 0), window=body, anchor="nw")
        body.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(win, width=e.width))
        self.body = body
        root.bind_all("<MouseWheel>", self.on_wheel)

        # ---------- HÀNG 1: NGUỒN + CẤU HÌNH ----------
        row1 = ttk.Frame(body)
        row1.pack(fill="x")

        c_src, src = self.make_card(row1, "1 · Nguồn dữ liệu")
        c_src.pack(side="left", fill="both", expand=True, padx=(0, 8))
        src.columnconfigure(1, weight=1)
        ttk.Label(src, text="HTML / URL", style="Card.TLabel").grid(row=0, column=0, sticky="w", pady=4)
        self.html_var = tk.StringVar(value="")
        ttk.Entry(src, textvariable=self.html_var, width=48).grid(row=0, column=1, sticky="ew", padx=8, pady=4)
        ttk.Button(src, text="Browse", style="Small.TButton", command=self.browse_html).grid(row=0, column=2, pady=4)
        ttk.Button(src, text="🔍 Quét form", style="Primary.TButton", command=self.scan_form_thread).grid(
            row=0, column=3, padx=(8, 0), pady=4)
        ttk.Label(src, text="Excel", style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=4)
        self.excel_var = tk.StringVar(value="")
        ttk.Entry(src, textvariable=self.excel_var, width=48).grid(row=1, column=1, sticky="ew", padx=8, pady=4)
        ttk.Button(src, text="Browse", style="Small.TButton", command=self.browse_excel).grid(row=1, column=2, pady=4)
        ttk.Button(src, text="📂 Load Excel", style="Primary.TButton", command=self.load_excel).grid(
            row=1, column=3, padx=(8, 0), pady=4)
        ttk.Label(src, text="Nhận file HTML trên máy hoặc đường dẫn http(s). Excel: sheet đầu tiên, hàng 1 là tiêu đề.",
                  style="Muted.TLabel").grid(row=2, column=0, columnspan=4, sticky="w", pady=(6, 0))

        c_cfg, cfg = self.make_card(row1, "2 · Cấu hình nhận diện")
        c_cfg.pack(side="left", fill="both", expand=True, padx=(8, 0))
        cfg.columnconfigure(1, weight=1)
        cfg.columnconfigure(5, weight=1)
        ttk.Label(cfg, text="Nút submit (CSS)", style="Card.TLabel").grid(row=0, column=0, sticky="w", pady=4)
        self.submit_var = tk.StringVar(value="")
        ttk.Entry(cfg, textvariable=self.submit_var, width=22).grid(row=0, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(cfg, text="🖱 Chọn", style="Small.TButton",
                   command=lambda: self.pick_thread("submit")).grid(row=0, column=2, pady=4)
        ttk.Label(cfg, text="Vùng báo lỗi (CSS)", style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=4)
        self.error_var = tk.StringVar(value="")
        ttk.Entry(cfg, textvariable=self.error_var, width=22).grid(row=1, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(cfg, text="🖱 Chọn", style="Small.TButton",
                   command=lambda: self.pick_thread("error")).grid(row=1, column=2, pady=4)
        ttk.Label(cfg, text="Vùng thành công (CSS)", style="Card.TLabel").grid(row=2, column=0, sticky="w", pady=4)
        self.msg_var = tk.StringVar(value="")
        ttk.Entry(cfg, textvariable=self.msg_var, width=22).grid(row=2, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(cfg, text="🖱 Chọn", style="Small.TButton",
                   command=lambda: self.pick_thread("msg")).grid(row=2, column=2, pady=4)

        ttk.Label(cfg, text="Format ngày Excel", style="Card.TLabel").grid(row=0, column=4, sticky="w", padx=(22, 0), pady=4)
        self.format_var = tk.StringVar(value="dd/mm/yyyy")
        ttk.Combobox(cfg, textvariable=self.format_var, values=list(DATE_FORMATS.keys()),
                     width=16, state="readonly").grid(row=0, column=5, sticky="ew", padx=6, pady=4)
        ttk.Label(cfg, text="Cột Expected", style="Card.TLabel").grid(row=1, column=4, sticky="w", padx=(22, 0), pady=4)
        self.expected_col_var = tk.StringVar(value="Expected")
        ttk.Entry(cfg, textvariable=self.expected_col_var, width=18).grid(row=1, column=5, sticky="ew", padx=6, pady=4)
        ttk.Label(cfg, text="Từ khoá thành công", style="Card.TLabel").grid(row=2, column=4, sticky="w", padx=(22, 0), pady=4)
        self.success_var = tk.StringVar(value="thành công, success, successful, đã đặt, đã đăng ký")
        ttk.Entry(cfg, textvariable=self.success_var, width=18).grid(row=2, column=5, sticky="ew", padx=6, pady=4)

        # ---------- THANH HÀNH ĐỘNG + THỐNG KÊ ----------
        bar = ttk.Frame(top_fixed)
        bar.pack(fill="x", pady=(0, 6))
        ttk.Button(bar, text="🤖 Tự động map", style="Ghost.TButton", command=self.auto_map).pack(side="left")
        ttk.Button(bar, text="▶  CHẠY TEST", style="Success.TButton", command=self.run_tests_thread).pack(
            side="left", padx=8)
        ttk.Button(bar, text="⏹ Dừng", style="Danger.TButton", command=self.stop_tests).pack(side="left")
        self.btn_open = ttk.Button(bar, text="📄 Mở báo cáo", style="Primary.TButton",
                                   command=self.open_report, state="disabled")
        self.btn_open.pack(side="left", padx=8)

        self.stat_labels = {}
        for key, text in (("ERROR", "ERROR  0"), ("REVIEW", "REVIEW  0"), ("FAIL", "FAIL  0"),
                          ("PASS", "PASS  0"), ("TOTAL", "Tổng  0")):
            bgc, fgc = STAT_COLORS[key]
            lb = tk.Label(bar, text=text, bg=bgc, fg=fgc, font=("Segoe UI", 10, "bold"), padx=14, pady=6)
            lb.pack(side="right", padx=(6, 0))
            self.stat_labels[key] = lb

        # ---------- HÀNG 2: FIELD + MAP ----------
        row2 = ttk.Frame(body)
        row2.pack(fill="both", expand=False)

        c_fields, fbody = self.make_card(row2, "3 · Field phát hiện trong web")
        c_fields.pack(side="left", fill="both", expand=True, padx=(0, 8))
        cols_f = ("ID/Name", "Name", "Loại", "Label")
        self.tree_fields = ttk.Treeview(fbody, columns=cols_f, show="headings", height=6)
        for c in cols_f:
            self.tree_fields.heading(c, text=c)
            self.tree_fields.column(c, width=110)
        self.tree_fields.column("Label", width=230)
        sbf = ttk.Scrollbar(fbody, orient="vertical", command=self.tree_fields.yview)
        self.tree_fields.configure(yscrollcommand=sbf.set)
        self.tree_fields.pack(side="left", fill="both", expand=True)
        sbf.pack(side="right", fill="y")

        c_map, mbody = self.make_card(row2, "4 · Map cột Excel → field web")
        c_map.pack(side="left", fill="both", expand=True, padx=(8, 0))
        cols_m = ("Cột Excel", "Map tới Field Web", "Ví dụ dữ liệu")
        mtree_wrap = ttk.Frame(mbody, style="Card.TFrame")
        mtree_wrap.pack(fill="both", expand=True)
        self.tree_map = ttk.Treeview(mtree_wrap, columns=cols_m, show="headings", height=6)
        for c in cols_m:
            self.tree_map.heading(c, text=c)
            self.tree_map.column(c, width=130)
        self.tree_map.column("Ví dụ dữ liệu", width=200)
        sbm = ttk.Scrollbar(mtree_wrap, orient="vertical", command=self.tree_map.yview)
        self.tree_map.configure(yscrollcommand=sbm.set)
        self.tree_map.pack(side="left", fill="both", expand=True)
        sbm.pack(side="right", fill="y")
        btn_frame = ttk.Frame(mbody, style="Card.TFrame")
        btn_frame.pack(fill="x", pady=(8, 0))
        ttk.Button(btn_frame, text="🤖 Tự động map", style="Small.TButton", command=self.auto_map).pack(side="left")
        ttk.Button(btn_frame, text="🗑 Xóa map", style="Small.TButton", command=self.clear_map).pack(side="left", padx=6)
        ttk.Label(btn_frame, text="Double-click một dòng để map / bỏ map thủ công",
                  style="Muted.TLabel").pack(side="left", padx=10)
        self.tree_map.bind("<Double-1>", self.edit_mapping)

        # ---------- HÀNG 3: KẾT QUẢ ----------
        c_res, rbody = self.make_card(body, "5 · Test data & kết quả")
        c_res.pack(fill="both", expand=True, pady=(10, 0))
        self.c_res = c_res
        cols_d = ("TC_ID", "Dữ liệu Excel (raw)", "Expected", "Status", "Lỗi / thông báo thực tế")
        self.tree_data = ttk.Treeview(rbody, columns=cols_d, show="headings", height=9)
        for c in cols_d:
            self.tree_data.heading(c, text=c)
            self.tree_data.column(c, width=100)
        self.tree_data.column("TC_ID", width=80, stretch=False)
        self.tree_data.column("Dữ liệu Excel (raw)", width=380)
        self.tree_data.column("Expected", width=150, stretch=False)
        self.tree_data.column("Status", width=110, stretch=False, anchor="center")
        self.tree_data.column("Lỗi / thông báo thực tế", width=360)
        self.tree_data.tag_configure("PASS", background="#ecfdf5", foreground="#166534")
        self.tree_data.tag_configure("FAIL", background="#fef2f2", foreground="#991b1b")
        self.tree_data.tag_configure("REVIEW", background="#fffbeb", foreground="#92400e")
        self.tree_data.tag_configure("ERROR", background="#f5f3ff", foreground="#5b21b6")
        sb = ttk.Scrollbar(rbody, orient="vertical", command=self.tree_data.yview)
        self.tree_data.configure(yscrollcommand=sb.set)
        self.tree_data.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

        # ---------- TIẾN TRÌNH + NHẬT KÝ ----------
        prog = ttk.Frame(top_fixed)
        prog.pack(fill="x", pady=(0, 8))
        self.progress = ttk.Progressbar(prog, mode="determinate", style="Modern.Horizontal.TProgressbar")
        self.progress.pack(side="left", fill="x", expand=True)
        self.progress_lbl = ttk.Label(prog, text="0 / 0", font=FONT_B)
        self.progress_lbl.pack(side="left", padx=(10, 0))

        self.log = tk.Text(body, height=6, wrap="word", bg="#0f172a", fg="#e2e8f0", insertbackground="white",
                           font=("Consolas", 9), relief="flat", padx=10, pady=8)
        self.log.pack(fill="x", pady=(10, 0))

        self.log_insert("Sẵn sàng. Quét form -> Load Excel -> Tự động map -> Chạy test.")
        self.log_insert("PASS/FAIL chỉ chấm khi bắt được lỗi hoặc thông báo; không bắt được gì -> REVIEW.")

    # ------------------------------------------------------------------ Giao diện: style & card
    def setup_style(self):
        st = ttk.Style()
        try:
            st.theme_use("clam")
        except Exception:
            pass
        st.configure(".", font=FONT, background=BG, foreground=TEXT)
        st.configure("TFrame", background=BG)
        st.configure("TLabel", background=BG, foreground=TEXT)
        st.configure("Card.TFrame", background=CARD)
        st.configure("Card.TLabel", background=CARD, foreground=TEXT)
        st.configure("Muted.TLabel", background=CARD, foreground=MUTED, font=("Segoe UI", 9))
        st.configure("CardTitle.TLabel", background=CARD, foreground=PRIMARY_DARK, font=("Segoe UI", 11, "bold"))

        for name in ("TEntry", "TCombobox"):
            st.configure(name, fieldbackground="white", foreground=TEXT, bordercolor="#cbd5e1",
                         lightcolor="#cbd5e1", darkcolor="#cbd5e1", padding=5)
            st.map(name, bordercolor=[("focus", PRIMARY)], lightcolor=[("focus", PRIMARY)],
                   darkcolor=[("focus", PRIMARY)])

        def button(name, bg, fg, hover, pad=(14, 8)):
            st.configure(name, background=bg, foreground=fg, borderwidth=0, focusthickness=0,
                         padding=pad, font=FONT_B)
            st.map(name, background=[("disabled", "#e2e8f0"), ("active", hover)],
                   foreground=[("disabled", "#94a3b8")])
        button("Primary.TButton", PRIMARY, "white", "#1d4ed8")
        button("Success.TButton", "#16a34a", "white", "#15803d", pad=(20, 8))
        button("Danger.TButton", "#dc2626", "white", "#b91c1c")
        button("Ghost.TButton", "#e2e8f0", TEXT, "#cbd5e1")
        button("Small.TButton", "#e2e8f0", TEXT, "#cbd5e1", pad=(10, 5))

        st.configure("Treeview", rowheight=28, background="white", fieldbackground="white",
                     foreground=TEXT, borderwidth=0, font=FONT)
        st.map("Treeview", background=[("selected", "#dbeafe")], foreground=[("selected", TEXT)])
        st.configure("Treeview.Heading", background="#1e293b", foreground="white", font=FONT_B,
                     padding=(8, 6), relief="flat", borderwidth=0)
        st.map("Treeview.Heading", background=[("active", "#334155")])
        st.configure("Vertical.TScrollbar", background="#cbd5e1", troughcolor=BG, bordercolor=BG,
                     arrowcolor="#64748b")
        st.configure("Modern.Horizontal.TProgressbar", troughcolor="#e2e8f0", background=PRIMARY,
                     bordercolor="#e2e8f0", lightcolor=PRIMARY, darkcolor=PRIMARY, thickness=12)

    def on_wheel(self, event):
        """Con lăn chuột cuộn cả trang; riêng bảng/nhật ký/danh sách xổ xuống tự cuộn riêng."""
        if isinstance(event.widget, (ttk.Treeview, tk.Text, tk.Listbox)):
            return
        self.canvas.yview_scroll(-1 if event.delta > 0 else 1, "units")

    def scroll_to_results(self):
        try:
            self.root.update_idletasks()
            total = max(1, self.body.winfo_height())
            self.canvas.yview_moveto(max(0.0, (self.c_res.winfo_y() - 10) / total))
        except Exception:
            pass

    def make_card(self, parent, title):
        """Khung trắng bo viền mảnh có tiêu đề. Trả về (khung ngoài, khung nội dung)."""
        outer = tk.Frame(parent, bg=CARD, highlightbackground=LINE, highlightthickness=1)
        inner = ttk.Frame(outer, style="Card.TFrame", padding=(14, 10))
        inner.pack(fill="both", expand=True)
        ttk.Label(inner, text=title, style="CardTitle.TLabel").pack(anchor="w", pady=(0, 8))
        content = ttk.Frame(inner, style="Card.TFrame")
        content.pack(fill="both", expand=True)
        return outer, content

    def update_stats(self, reset=False):
        res = [] if reset else self.results
        cnt = {k: sum(1 for r in res if r["status"] == k) for k in ("PASS", "FAIL", "REVIEW", "ERROR")}
        total = len(self.test_data)
        self.stat_labels["TOTAL"].config(text=f"Tổng  {total}")
        for k, v in cnt.items():
            self.stat_labels[k].config(text=f"{k}  {v}")
        self.progress_lbl.config(text=f"{len(res)} / {total}")

    # ------------------------------------------------------------------ UI helpers
    def log_insert(self, msg):
        line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n"

        def _do():
            self.log.insert("end", line)
            self.log.see("end")

        if threading.current_thread() is threading.main_thread():
            _do()
        else:
            self.root.after(0, _do)

    def ui_error(self, title, msg):
        self.root.after(0, lambda: messagebox.showerror(title, msg))

    def ui_warn(self, title, msg):
        self.root.after(0, lambda: messagebox.showwarning(title, msg))

    def browse_html(self):
        p = filedialog.askopenfilename(filetypes=[("HTML", "*.html *.htm")])
        if p:
            self.html_var.set(p)

    def browse_excel(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if p:
            self.excel_var.set(p)

    def success_kws(self):
        return [norm(x) for x in re.split(r"[,;]", self.success_var.get()) if norm(x)]

    def make_driver(self, headless=False):
        options = webdriver.ChromeOptions()
        if headless:
            options.add_argument("--headless=new")
        # Giữ alert lại để tool tự đọc nội dung (mặc định Selenium sẽ đóng alert)
        options.set_capability("unhandledPromptBehavior", "ignore")
        return webdriver.Chrome(options=options)

    # ------------------------------------------------------------------ Chọn phần tử bằng chuột
    def pick_thread(self, target):
        threading.Thread(target=self.pick_element, args=(target,), daemon=True).start()

    def pick_element(self, target):
        """Mở Chrome, người dùng GIỮ CTRL + CLICK vào phần tử -> lưu selector CSS.
        target: 'submit' (nút gửi), 'error' (vùng báo lỗi - ưu tiên selector dùng chung cho cả nhóm),
        'msg' (vùng thông báo thành công)."""
        src = self.html_var.get().strip()
        if not source_exists(src):
            self.ui_error("Lỗi", f"Không thấy file HTML / URL:\n{src}")
            return
        names = {"submit": "nút Đăng ký", "error": "vùng báo lỗi", "msg": "vùng thông báo thành công"}
        hint = {
            "submit": "Giữ Ctrl + click vào nút Đăng ký.",
            "error": "Bấm Đăng ký khi để trống để lỗi hiện ra, rồi giữ Ctrl + click vào một dòng lỗi.",
            "msg": "Điền dữ liệu hợp lệ, bấm Đăng ký để hiện thông báo, rồi giữ Ctrl + click vào thông báo.",
        }
        self.log_insert(f"Mở trình duyệt để chọn {names[target]}. {hint[target]} (tối đa 3 phút)")
        driver = None
        try:
            driver = self.make_driver(headless=False)
            driver.get(to_url(src))
            WebDriverWait(driver, 10).until(
                lambda d: d.execute_script("return document.readyState") == "complete")
            group = (target == "error")
            picked, deadline = None, time.time() + 180
            while time.time() < deadline and not picked:
                try:
                    driver.execute_script(JS_PICK, group)           # nạp lại nếu trang vừa chuyển
                    picked = driver.execute_script("return window.__picked || null;")
                except NoAlertPresentException:
                    pass
                except Exception as e:
                    if "alert" in str(e).lower():
                        self.read_alert(driver)                      # đóng alert để tiếp tục chọn
                    else:
                        break                                        # trình duyệt đã bị đóng
                time.sleep(0.3)
            if not picked:
                self.log_insert("Không chọn được phần tử (hết giờ hoặc đã đóng trình duyệt).")
                return
            try:
                count = len(driver.find_elements(By.CSS_SELECTOR, picked))
            except Exception:
                count = "?"
            var = {"submit": self.submit_var, "error": self.error_var, "msg": self.msg_var}[target]
            self.root.after(0, lambda: var.set(picked))
            self.log_insert(f"Đã chọn {names[target]}: {picked}  (khớp {count} phần tử)")
        except Exception as e:
            self.log_insert(f"Lỗi khi chọn phần tử: {e}")
            self.ui_error("Lỗi chọn phần tử", str(e))
        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass

    # ------------------------------------------------------------------ Quét form
    def scan_form_thread(self):
        threading.Thread(target=self.scan_form, daemon=True).start()

    def scan_form(self):
        src = self.html_var.get().strip()
        if not source_exists(src):
            self.ui_error("Lỗi", f"Không thấy file HTML / URL:\n{src}")
            return
        self.log_insert(f"Đang quét form: {src}")
        driver = None
        try:
            driver = self.make_driver(headless=True)
            driver.get(to_url(src))
            WebDriverWait(driver, 10).until(
                lambda d: d.execute_script("return document.readyState") == "complete")
            time.sleep(0.5)

            fields = []
            for el in driver.find_elements(By.CSS_SELECTOR, "input, select, textarea"):
                try:
                    fid = el.get_attribute("id") or ""
                    fname = el.get_attribute("name") or ""
                    ftype = (el.get_attribute("type") or el.tag_name).lower()
                    if ftype in ("submit", "button", "hidden", "file", "reset", "image"):
                        continue
                    if not fid and not fname:
                        continue
                    try:
                        label_text = driver.execute_script(JS_LABEL, el) or ""
                    except Exception:
                        label_text = ""
                    # radio: gộp theo name (cả nhóm là 1 field)
                    key = fname if (ftype == "radio" and fname) else (fid or fname)
                    fields.append({
                        "id": fid, "name": fname, "type": ftype, "tag": el.tag_name.lower(),
                        "label": label_text[:50], "key": key,
                    })
                except Exception:
                    continue

            seen, uniq = set(), []
            for f in fields:
                if f["key"] and f["key"] not in seen:
                    seen.add(f["key"])
                    uniq.append(f)

            self.web_fields = uniq
            self.root.after(0, self.refresh_fields_tree)
            self.log_insert(f"Quét xong: {len(uniq)} field: {', '.join(f['key'] for f in uniq)}")
        except Exception as e:
            self.log_insert(f"Lỗi quét form: {e}")
            self.ui_error("Lỗi quét", str(e))
        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass

    def refresh_fields_tree(self):
        self.tree_fields.delete(*self.tree_fields.get_children())
        for f in self.web_fields:
            self.tree_fields.insert("", "end", values=(f["id"] or f["key"], f["name"], f["type"], f["label"]))

    # ------------------------------------------------------------------ Excel
    def load_excel(self):
        path = self.excel_var.get().strip()
        if not os.path.exists(path):
            messagebox.showerror("Lỗi", f"Không thấy Excel:\n{path}")
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active

            headers, used = [], {}
            for i, c in enumerate(ws[1], 1):
                h = str(c.value).strip() if c.value is not None and str(c.value).strip() else f"Col{i}"
                if h in used:
                    used[h] += 1
                    h = f"{h}_{used[h]}"
                else:
                    used[h] = 1
                headers.append(h)
            self.excel_headers = headers
            self.log_insert(f"Excel headers: {headers}")

            self.test_data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not any(v is not None and str(v).strip() for v in row):
                    continue
                row_dict = {h: (row[i] if i < len(row) else None) for i, h in enumerate(headers)}
                tc_id = str(row[0]).strip() if row[0] is not None and str(row[0]).strip() \
                    else f"TC{len(self.test_data) + 1}"
                self.test_data.append((tc_id, row_dict, row))

            self.refresh_data_tree()
            self.log_insert(f"Load {len(self.test_data)} test case từ Excel")

            if not self.mappings and self.web_fields:
                self.auto_map()
            else:
                self.refresh_map_tree()
        except Exception as e:
            messagebox.showerror("Lỗi Excel", str(e))

    def get_expected(self, row_dict):
        name = self.expected_col_var.get().strip()
        if not name:
            return ""
        n = norm(name)
        match = None
        for h in self.excel_headers:
            if norm(h) == n:
                match = h
                break
        if match is None and len(n) >= 3:
            for h in self.excel_headers:
                if n in norm(h):
                    match = h
                    break
        if match is None:
            return ""
        v = clean_cell(row_dict.get(match))
        return to_text(v) if v != "" else ""

    def refresh_data_tree(self):
        self.tree_data.delete(*self.tree_data.get_children())
        for i, (tc_id, row_dict, _raw) in enumerate(self.test_data):
            raw_str = ", ".join(f"{k}={v}" for k, v in row_dict.items() if v not in (None, ""))[:120]
            self.tree_data.insert("", "end", iid=str(i),
                                  values=(tc_id, raw_str, self.get_expected(row_dict), "Ready", ""))
        self.update_stats(reset=True)

    # ------------------------------------------------------------------ Map
    def is_non_input_header(self, header):
        h = norm(header)
        if h in NON_INPUT_HEADERS:
            return True
        if any(h.startswith(p + " ") for p in ("expected", "actual", "ghi chu")):
            return True
        e = norm(self.expected_col_var.get())
        return bool(e) and h == e

    @staticmethod
    def _match_score(h, h_ns, f):
        k = norm(f["key"]).replace(" ", "")
        lab = norm(f["label"])
        lab_ns = lab.replace(" ", "")
        if h_ns and (h_ns == k or (lab and h == lab)):
            return 100
        best = 0
        for s in (k, lab_ns):
            if len(s) >= 3 and len(h_ns) >= 3 and (s in h_ns or h_ns in s):
                best = max(best, min(len(s), len(h_ns)))
        return best

    def auto_map(self):
        if not self.web_fields or not self.excel_headers:
            messagebox.showwarning("Chưa đủ dữ liệu", "Hãy Quét Form Web và Load Excel trước")
            return

        self.mappings = {}
        used = set()
        pending = []

        # Lượt 1: khớp tên / label (chọn điểm cao nhất, mỗi field chỉ map 1 lần)
        for idx, header in enumerate(self.excel_headers):
            if self.is_non_input_header(header):
                continue
            h = norm(header)
            h_ns = h.replace(" ", "")
            best, best_score = None, 0
            for f in self.web_fields:
                if f["key"] in used:
                    continue
                sc = self._match_score(h, h_ns, f)
                if sc > best_score:
                    best, best_score = f, sc
            if best:
                self.mappings[idx] = best["key"]
                used.add(best["key"])
            else:
                pending.append(idx)

        # Lượt 2: từ đồng nghĩa theo loại field (không phụ thuộc loại web cụ thể)
        for idx in pending:
            h = norm(self.excel_headers[idx])
            for h_kws, f_kws, types, tags in SYNONYM_GROUPS:
                if not any(kw_hit(h, kw) for kw in h_kws):
                    continue
                cand = None
                for f in self.web_fields:
                    if f["key"] in used:
                        continue
                    k = norm(f["key"]).replace(" ", "")
                    lab = norm(f["label"]).replace(" ", "")
                    if f["type"] in types or f["tag"] in tags or any(w in k or w in lab for w in f_kws):
                        cand = f
                        break
                if cand:
                    self.mappings[idx] = cand["key"]
                    used.add(cand["key"])
                    break

        self.refresh_map_tree()
        unmapped = [self.excel_headers[i] for i in range(len(self.excel_headers))
                    if i not in self.mappings and not self.is_non_input_header(self.excel_headers[i])]
        unused = [f["key"] for f in self.web_fields if f["key"] not in used]
        self.log_insert(f"Tự động map: {len(self.mappings)} cột. "
                        f"Chưa map (Excel): {unmapped or 'không'} | Field web chưa dùng: {unused or 'không'}")

    def refresh_map_tree(self):
        self.tree_map.delete(*self.tree_map.get_children())
        for idx, header in enumerate(self.excel_headers):
            field = self.mappings.get(idx, "(bỏ qua)")
            example = ""
            if self.test_data:
                example = str(self.test_data[0][1].get(header, ""))[:30]
            self.tree_map.insert("", "end", iid=str(idx), values=(header, field, example))

    def clear_map(self):
        self.mappings = {}
        self.refresh_map_tree()

    def edit_mapping(self, _event):
        sel = self.tree_map.selection()
        if not sel or not self.web_fields:
            return
        idx = int(sel[0])
        header = self.excel_headers[idx]
        options = ["(bỏ qua)"] + [f["key"] for f in self.web_fields]

        popup = tk.Toplevel(self.root)
        popup.title("Chọn Field Web")
        popup.geometry("320x170")
        ttk.Label(popup, text=f"Cột Excel: {header}").pack(pady=5)
        var = tk.StringVar(value=self.mappings.get(idx, "(bỏ qua)"))
        ttk.Combobox(popup, textvariable=var, values=options, state="readonly", width=32).pack(pady=10)

        def save():
            if var.get() == "(bỏ qua)":
                self.mappings.pop(idx, None)
            else:
                self.mappings[idx] = var.get()
            self.refresh_map_tree()
            popup.destroy()

        ttk.Button(popup, text="Lưu", command=save).pack(pady=10)

    # ------------------------------------------------------------------ Điền dữ liệu
    @staticmethod
    def click(driver, el):
        try:
            el.click()
        except Exception:
            driver.execute_script("arguments[0].click()", el)

    @staticmethod
    def select_option(sel, s):
        for fn in (sel.select_by_visible_text, sel.select_by_value):
            try:
                fn(s)
                return True
            except NoSuchElementException:
                pass
        ns = norm(s)
        for i, o in enumerate(sel.options):
            if norm(o.text) == ns or norm(o.get_attribute("value") or "") == ns:
                sel.select_by_index(i)
                return True
        if ns:
            for i, o in enumerate(sel.options):
                if ns in norm(o.text):
                    sel.select_by_index(i)
                    return True
        return False

    def fill_field(self, driver, key, raw_value):
        """Điền 1 field. Trả danh sách ghi chú/cảnh báo."""
        notes = []
        val = clean_cell(raw_value)
        els = driver.find_elements(By.ID, key) or driver.find_elements(By.NAME, key)
        if not els:
            return [f"Không thấy field '{key}' trên web"]
        el = els[0]
        tag = el.tag_name.lower()
        ftype = (el.get_attribute("type") or tag).lower()
        s = to_text(val) if val != "" else ""

        if tag == "select":
            if s and not self.select_option(Select(el), s):
                notes.append(f"{key}: không có option '{s}'")

        elif ftype == "radio":
            name = el.get_attribute("name")
            group = [r for r in driver.find_elements(By.NAME, name)
                     if (r.get_attribute("type") or "").lower() == "radio"] if name else [el]
            if s:
                ns = norm(s)
                target = None
                for r in group:
                    lab = norm(driver.execute_script(JS_LABEL, r))
                    if norm(r.get_attribute("value") or "") == ns or lab == ns:
                        target = r
                        break
                if target is None:
                    for r in group:
                        if ns in norm(driver.execute_script(JS_LABEL, r)):
                            target = r
                            break
                if target is not None:
                    self.click(driver, target)
                else:
                    notes.append(f"{key}: không có lựa chọn '{s}'")

        elif ftype == "checkbox":
            want = norm(s) in TRUE_SET
            if el.is_selected() != want:
                self.click(driver, el)

        elif ftype == "date":
            iso, note = parse_date_flexible(val, self.format_var.get()) if val != "" else ("", "")
            if note:
                notes.append(f"{key}: {note}")
            driver.execute_script(JS_SET, el, iso)
            if iso and el.get_attribute("value") != iso:
                notes.append(f"{key}: trình duyệt không nhận ngày '{iso}'")

        elif ftype in JS_ONLY_TYPES:
            driver.execute_script(JS_SET, el, s)

        else:  # text, number, email, tel, password, textarea...
            try:
                el.clear()
                if s:
                    el.send_keys(s)
            except Exception:
                driver.execute_script(JS_SET, el, s)
            if s and not (el.get_attribute("value") or ""):
                notes.append(f"{key}: trình duyệt không nhận giá trị '{s}' (field rỗng)")
        return notes

    # ------------------------------------------------------------------ Submit & đọc kết quả
    def click_submit(self, driver):
        candidates = []
        custom = self.submit_var.get().strip()
        if custom:
            candidates.append(custom)
        candidates += [
            "#submit", "form button[type='submit']", "form input[type='submit']",
            "button[type='submit']", "input[type='submit']",
            "form button:not([type='button']):not([type='reset'])",
        ]
        for css in candidates:
            try:
                for el in driver.find_elements(By.CSS_SELECTOR, css):
                    if el.is_displayed() and el.is_enabled():
                        self.click(driver, el)
                        return ""
            except Exception:
                continue
        try:
            ok = driver.execute_script(
                "var f=document.querySelector('form'); if(!f) return false;"
                "if(f.requestSubmit) f.requestSubmit(); else f.submit(); return true;")
            if ok:
                return "Không thấy nút submit, đã submit form bằng JS"
        except Exception:
            pass
        return "KHÔNG tìm thấy nút submit"

    @staticmethod
    def read_alert(driver):
        try:
            al = driver.switch_to.alert
            text = (al.text or "").strip() or "(alert)"
            al.accept()
            return text
        except NoAlertPresentException:
            return ""
        except Exception:
            return ""

    def error_sel(self):
        return self.error_var.get().strip() or ERROR_SELECTOR

    def msg_sel(self):
        return self.msg_var.get().strip() or MSG_SELECTOR

    def read_dom_state(self, driver):
        """Trả (danh sách lỗi 'id: text', message text)."""
        entries = []
        try:
            for el in driver.find_elements(By.CSS_SELECTOR, self.error_sel()):
                try:
                    txt = el.text.strip()
                    if txt and len(txt) > 2:
                        eid = el.get_attribute("id") or el.get_attribute("class") or ""
                        entries.append((eid, txt))
                except Exception:
                    continue
        except Exception:
            pass
        # bỏ element cha chứa lại text của element con
        keep = [e for e in entries
                if not any(o is not e and len(o[1]) < len(e[1]) and o[1] in e[1] for o in entries)]
        errors, seen = [], set()
        for eid, txt in keep:
            item = f"{eid}: {txt}" if eid else txt
            if item not in seen:
                seen.add(item)
                errors.append(item)

        msg = ""
        try:
            for el in driver.find_elements(By.CSS_SELECTOR, self.msg_sel()):
                t = el.text.strip()
                if t:
                    msg = t
                    break
        except Exception:
            pass
        return errors, msg

    def wait_for_result(self, driver, pre_invalid, base_errors, base_msg, timeout=3.0):
        deadline = time.time() + (1.0 if pre_invalid else timeout)
        alert_text, errors, msg, found_at = "", [], "", None
        while time.time() < deadline:
            a = self.read_alert(driver)
            if a:
                alert_text = f"{alert_text} | {a}" if alert_text else a
            e, m = self.read_dom_state(driver)
            errors = [x for x in e if x not in base_errors]
            msg = m if m != base_msg else ""
            if alert_text or errors or msg:
                if found_at is None:
                    found_at = time.time()
                if time.time() - found_at >= 0.3:   # chờ thêm chút để gom đủ lỗi
                    break
            time.sleep(0.1)
        validation = []
        if not (alert_text or errors or msg) and pre_invalid:
            validation = [f"HTML5: {x}" for x in pre_invalid]
        return {"alert": alert_text, "errors": errors, "msg": msg, "validation": validation}

    def classify(self, st):
        text = " | ".join(t for t in (st["alert"], st["msg"]) if t)
        errs = st["errors"] + st["validation"]
        if errs:
            return "ERROR", " | ".join(errs)
        if text:
            tn = norm(text)
            if any(kw_hit(tn, k) for k in FAIL_KWS):
                return "ERROR", text
            if any(kw_hit(tn, k) for k in self.success_kws()):
                return "SUCCESS", text
            return "UNKNOWN", text
        return "UNKNOWN", ""

    @staticmethod
    def judge(expected, should_pass, exp_msg, actual, err_str):
        if expected == "" or should_pass is None:
            return "REVIEW", "Không có Expected"
        if actual == "UNKNOWN":
            return "REVIEW", "Không bắt được lỗi/thông báo nào -> cần xem tay"
        if (actual == "SUCCESS") != should_pass:
            return "FAIL", ("Mong đợi thành công nhưng web báo lỗi" if should_pass
                            else "Mong đợi lỗi nhưng web báo thành công")
        if not should_pass and exp_msg and not msg_matches(exp_msg, err_str):
            return "FAIL", f"Sai nội dung lỗi (mong đợi: '{exp_msg}')"
        return "PASS", ""

    # ------------------------------------------------------------------ Chạy test
    def stop_tests(self):
        if self.running:
            self.stop_event.set()
            self.log_insert("Đã yêu cầu dừng sau case hiện tại...")

    def run_tests_thread(self):
        if self.running:
            messagebox.showinfo("Đang chạy", "Test đang chạy, bấm Dừng nếu muốn huỷ.")
            return
        if not self.test_data:
            messagebox.showwarning("Chưa load", "Load Excel trước")
            return
        if not self.web_fields:
            messagebox.showwarning("Chưa quét", "Quét Form Web trước")
            return
        if not self.mappings:
            messagebox.showwarning("Chưa map", "Map cột Excel -> Field Web trước (bấm Tự Động Map)")
            return
        if not source_exists(self.html_var.get()):
            messagebox.showerror("Lỗi", "Không thấy file HTML / URL")
            return
        self.refresh_data_tree()
        self.scroll_to_results()
        threading.Thread(target=self.run_tests, daemon=True).start()

    def run_one(self, driver, url, tc_id, row_dict):
        driver.get(url)   # tải lại sạch (refresh có thể khôi phục giá trị cũ của form)
        WebDriverWait(driver, 10).until(
            lambda d: d.execute_script("return document.readyState") == "complete")
        self.read_alert(driver)

        notes = []
        for excel_idx, key in sorted(self.mappings.items()):
            if excel_idx >= len(self.excel_headers):
                continue
            header = self.excel_headers[excel_idx]
            try:
                notes += self.fill_field(driver, key, row_dict.get(header))
            except Exception as e:
                notes.append(f"{key}: lỗi khi điền ({type(e).__name__}: {e})")

        try:
            pre_invalid = driver.execute_script(JS_INVALID) or []
        except Exception:
            pre_invalid = []
        base_errors, base_msg = self.read_dom_state(driver)

        submit_note = self.click_submit(driver)
        if submit_note:
            notes.append(submit_note)

        st = self.wait_for_result(driver, pre_invalid, base_errors, base_msg)
        actual, err_str = self.classify(st)

        expected = self.get_expected(row_dict)
        should_pass, exp_msg = parse_expected(expected, self.success_kws())
        status, judge_note = self.judge(expected, should_pass, exp_msg, actual, err_str)
        if judge_note:
            notes.append(judge_note)

        shot = ""
        if status != "PASS":
            try:
                shot = driver.get_screenshot_as_base64()
            except Exception:
                shot = ""
        raw_data = ", ".join(f"{k}={v}" for k, v in row_dict.items() if v not in (None, ""))[:200]
        return {
            "tc_id": tc_id, "raw": raw_data, "expected": expected, "actual": actual,
            "status": status, "error": err_str, "note": " ; ".join(notes),
            "data": self.data_pairs(row_dict), "scenario": self.scenario_of(row_dict), "shot": shot,
        }

    def run_tests(self):
        self.running = True
        self.stop_event.clear()
        self.results = []
        self.run_started = time.time()
        total = len(self.test_data)
        self.root.after(0, lambda: self.progress.configure(maximum=total, value=0))
        driver = None
        try:
            url = to_url(self.html_var.get())
            driver = self.make_driver(headless=False)
            for idx, (tc_id, row_dict, _raw) in enumerate(self.test_data):
                if self.stop_event.is_set():
                    self.log_insert("Đã dừng theo yêu cầu.")
                    break
                self.log_insert(f"Chạy {tc_id}...")
                t0 = time.time()
                try:
                    res = self.run_one(driver, url, tc_id, row_dict)
                except Exception as e:
                    res = {
                        "tc_id": tc_id,
                        "raw": ", ".join(f"{k}={v}" for k, v in row_dict.items() if v not in (None, ""))[:200],
                        "expected": self.get_expected(row_dict), "actual": "ERROR",
                        "status": "ERROR", "error": f"{type(e).__name__}: {e}", "note": "",
                    }
                res.setdefault("data", self.data_pairs(row_dict))
                res.setdefault("scenario", self.scenario_of(row_dict))
                res.setdefault("shot", "")
                res["duration"] = round(time.time() - t0, 2)
                self.results.append(res)
                self.root.after(0, lambda i=idx, r=res: self.update_data_row(i, r))
                self.root.after(0, lambda v=idx + 1: self.progress.configure(value=v))
        except Exception as e:
            self.log_insert(f"Lỗi chạy test: {e}")
            self.ui_error("Lỗi chạy test", str(e))
        finally:
            if driver:
                try:
                    driver.quit()
                except Exception:
                    pass
            self.running = False
            self.log_insert(f"Hoàn thành {len(self.results)}/{total} case")
            if self.results:
                self.export_report()

    def update_data_row(self, idx, res):
        try:
            iid = str(idx)
            self.tree_data.set(iid, "Status", STATUS_ICON.get(res["status"], res["status"]))
            self.tree_data.set(iid, "Lỗi / thông báo thực tế", (res["error"] or res["note"])[:100])
            self.tree_data.item(iid, tags=(res["status"],))
            self.tree_data.see(iid)
            self.update_stats()
        except Exception:
            pass

    # ------------------------------------------------------------------ Report
    def data_pairs(self, row_dict):
        """Danh sách [tên cột, giá trị] của các cột đã map (dữ liệu nhập)."""
        out = []
        for idx, key in sorted(self.mappings.items()):
            if idx < len(self.excel_headers):
                h = self.excel_headers[idx]
                v = row_dict.get(h)
                out.append([h, "" if v is None else str(v)])
        return out

    def scenario_of(self, row_dict):
        """Lấy mô tả test case nếu Excel có cột Test Scenario / Mô tả / Test case Name."""
        names = {"test scenario", "scenario", "mo ta", "description", "test case name", "ten test case"}
        for h, v in row_dict.items():
            if norm(h) in names and v not in (None, ""):
                return str(v)
        return ""

    def export_report(self):
        import json
        src = self.html_var.get().strip()
        base_dir = os.path.dirname(src) if os.path.exists(src) else os.getcwd()
        report_path = os.path.join(base_dir or ".", f"report_universal_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html")
        cnt = {k: sum(1 for r in self.results if r["status"] == k) for k in ("PASS", "FAIL", "REVIEW", "ERROR")}
        meta = {
            "time": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "seconds": round(time.time() - getattr(self, "run_started", time.time()), 1),
            "web": src, "excel": self.excel_var.get(), "date_format": self.format_var.get(),
            "mapping": [[self.excel_headers[i], k] for i, k in sorted(self.mappings.items())
                        if i < len(self.excel_headers)],
            "selectors": {"submit": self.submit_var.get() or "(tự tìm)",
                          "error": self.error_var.get() or "(mặc định)",
                          "msg": self.msg_var.get() or "(mặc định)"},
            "total": len(self.results), "counts": cnt,
        }
        safe = lambda o: json.dumps(o, ensure_ascii=False).replace("</", "<\\/")
        page = REPORT_TEMPLATE.replace("__META__", safe(meta)).replace("__RESULTS__", safe(self.results))

        with open(report_path, "w", encoding="utf-8") as f:
            f.write(page)

        self.report_path = report_path
        self.log_insert(f"Đã xuất report: {report_path}")
        self.root.after(0, lambda: self.btn_open.config(state="normal"))
        webbrowser.open(Path(report_path).resolve().as_uri())

    def open_report(self):
        if self.report_path and os.path.exists(self.report_path):
            webbrowser.open(Path(self.report_path).resolve().as_uri())


if __name__ == "__main__":
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)   # chữ sắc nét trên màn hình độ phân giải cao (Windows)
    except Exception:
        pass
    root = tk.Tk()
    app = UniversalTesterGUI(root)
    root.mainloop()
